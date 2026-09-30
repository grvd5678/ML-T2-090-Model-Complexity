from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

print("Drafting the official Stage 1 Research Brief... ✍️✨")

doc = Document()

# Traditional Title Formatting
title = doc.add_heading('Stage 1 Research Brief: Problem Understanding & Initial Research', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
doc.add_paragraph('Project: ML-T2-090 | Intern ID: [Your ID] | Domain: Credit Risk').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
doc.add_paragraph('') # Spacer

# 1. Problem Understanding
doc.add_heading('1. Problem Understanding', level=1)
doc.add_paragraph("The real-world problem is determining whether added algorithmic complexity genuinely provides a justifiable benefit over simple, traditional baselines. In highly regulated environments like banking, machine learning models must evaluate credit risk safely and transparently. We need to understand if trendy, complex models actually solve this problem better than classic, easily explainable statistical methods.")

# 2. Problem Formulation
doc.add_heading('2. Problem Formulation', level=1)
doc.add_paragraph("Task: Binary Classification.\nInputs: Applicant financial history, employment status, and loan details (mixed categorical and numerical data).\nOutput/Target: A binary prediction of 'Good' or 'Bad' credit risk.\nUnit of Prediction: A single loan applicant.\nConstraints: The solution must rely exclusively on open-source tools and legally usable public data. The final model must maintain high explainability for regulatory compliance.")

# 3. Use Case
doc.add_heading('3. Use Case', level=1)
doc.add_paragraph("The intended users are traditional banking institutions and loan officers. This matters because falsely approving a high-risk applicant costs the bank significantly more than rejecting a safe applicant. Furthermore, the institution legally requires a transparent model to explain any loan denials to customers.")

# 4. Initial Data Strategy
doc.add_heading('4. Initial Data Strategy', level=1)
doc.add_paragraph("I plan to use the historically significant OpenML German Credit dataset. It is public, completely free to use, and contains 1000 records of traditional financial features. A known concern is the realistic class imbalance (70% good credit, 30% bad credit), which will require careful evaluation to prevent the model from blindly guessing the majority class.")

# 5. Existing Solutions
doc.add_heading('5. Existing Solutions', level=1)
doc.add_paragraph("Historically, the financial sector has relied heavily on Logistic Regression because of its mathematical stability and interpretability. Recently, there has been a push to use complex ensembles (like Random Forests or Gradient Boosting), but these often struggle with the strict explainability requirements of the industry.")

# 6. Literature / Technical Research
doc.add_heading('6. Literature / Technical Research', level=1)
doc.add_paragraph("Initial research into the scikit-learn documentation confirms that handling the mixed data types in financial datasets requires strict, traditional preprocessing. Numerical features must be normalized (StandardScaler) to prevent magnitude bias, and categorical text must be converted to binary arrays (OneHotEncoder) while avoiding the dummy variable trap.")

# 7. Candidate Approaches
doc.add_heading('7. Candidate Approaches', level=1)
doc.add_paragraph("Approach A (The Baseline): Logistic Regression. It fits the traditional banking need for explainability and speed, though it might struggle if the data has highly non-linear relationships.\nApproach B (The Complex Alternative): Random Forest Classifier. It can handle complex data relationships easily, but it acts as a 'black box' and may mistakenly overfit the majority class in an imbalanced dataset.")

# 8. Evaluation Strategy
doc.add_heading('8. Evaluation Strategy', level=1)
doc.add_paragraph("Due to the class imbalance, overall accuracy is a misleading metric. The primary metric will be Recall for the minority class ('Bad Credit'), as failing to identify a bad loan carries the highest business cost. Engineering inference time will also be recorded to measure computational cost.")

# 9. Initial Methodology
doc.add_heading('9. Initial Methodology', level=1)
doc.add_paragraph("1. Fetch the OpenML dataset.\n2. Apply traditional preprocessing (scaling and encoding) to maintain data integrity.\n3. Perform a standard 80/20 train-test split.\n4. Train the baseline Logistic Regression model.\n5. Train the complex Random Forest model.\n6. Compare both models strictly on business-critical recall and execution time.")

# 10. Open Questions
doc.add_heading('10. Open Questions', level=1)
doc.add_paragraph("Will the complex Random Forest actually over-optimize for the 'Good Credit' majority class? Does the computational cost of the complex model justify its performance, or will the simple baseline prove to be more protective of the bank's assets?")

# 11. References
doc.add_heading('11. References', level=1)
doc.add_paragraph("- OpenML German Credit Dataset (credit-g, version 1)\n- Scikit-Learn Official Documentation (Linear Models & Ensemble Methods)\n- Learn Depth Academy Track 2 Project 01 Orientation Guide")

# Save Document
doc.save('docs/Stage1_Research_Brief.docx')
print("Flawless Stage 1 Brief successfully saved to docs/Stage1_Research_Brief.docx! 💅")

